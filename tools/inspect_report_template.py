# -*- coding: utf-8 -*-
"""Dump the structure of the training-report template so we can match it."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

TEMPLATE = Path(r"C:\Users\y\Downloads\实训报告模板_视觉智能体综合实训.docx")
OUT = Path(__file__).resolve().parent / "template_structure.txt"


def iter_block_items(parent):
    """Yield Paragraph and Table in document order."""
    from docx.oxml.ns import qn
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def main() -> None:
    doc = Document(str(TEMPLATE))
    lines: list[str] = []

    def add(s: str = "") -> None:
        lines.append(s)

    add("=" * 90)
    add(f"模板：{TEMPLATE.name}")
    add("=" * 90)

    # 页面设置
    for i, sec in enumerate(doc.sections):
        add(f"\n[节 {i}] 纸张 {sec.page_width}x{sec.page_height} "
            f"边距 上{sec.top_margin} 下{sec.bottom_margin} "
            f"左{sec.left_margin} 右{sec.right_margin}")
        if sec.header and sec.header.paragraphs:
            add(f"  页眉：{' | '.join(p.text for p in sec.header.paragraphs if p.text)}")
        if sec.footer and sec.footer.paragraphs:
            add(f"  页脚：{' | '.join(p.text for p in sec.footer.paragraphs if p.text)}")

    add("\n" + "=" * 90)
    add("正文结构（按文档顺序）")
    add("=" * 90)

    n_p = n_t = 0
    for blk in iter_block_items(doc):
        if isinstance(blk, Paragraph):
            n_p += 1
            txt = blk.text.strip()
            style = blk.style.name if blk.style is not None else "?"
            if not txt:
                continue
            # 抓字体信息（第一个 run）
            fname = fsize = None
            bold = None
            if blk.runs:
                r = blk.runs[0]
                fname = r.font.name
                fsize = r.font.size.pt if r.font.size else None
                bold = r.font.bold
            add(f"[P{n_p:03d}] <{style}> "
                f"{'B' if bold else ' '} {fname or '-'}/{fsize or '-'}  {txt[:120]}")
        else:
            n_t += 1
            rows = len(blk.rows)
            cols = len(blk.columns)
            add(f"\n[T{n_t:02d}] 表格 {rows} 行 × {cols} 列")
            for ri, row in enumerate(blk.rows[:12]):
                cells = [c.text.strip().replace("\n", " / ")[:40] for c in row.cells]
                add(f"     r{ri}: {cells}")
            if rows > 12:
                add(f"     ... 还有 {rows - 12} 行")
            add("")

    add("=" * 90)
    add(f"合计：段落 {n_p} 个，表格 {n_t} 个")

    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"[写出] {OUT}")
    print(f"段落 {n_p}，表格 {n_t}")


if __name__ == "__main__":
    main()
